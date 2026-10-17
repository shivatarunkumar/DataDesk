"""Convertor: turn the sheets of an Excel workbook (.xlsx) into CSV files in My files.

Each chosen sheet becomes one CSV, written with the chosen delimiter and saved as a normal
file (version 1) in the folder the person picks. Cells are written the way the Cloud Run
file processor reads Excel (file-processor-service/app/converter.py): a formula gives its
last calculated value, a date at midnight is a plain ISO date, 12.0 is 12. Fully empty
rows and trailing empty columns are dropped.

Only .xlsx is read (openpyxl, already used for Excel everywhere else); old .xls files are
refused with what to do instead.
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.models.user import User
from app.models.workspace import WorkspaceFile
from app.services import workspace as ws
from app.services.errors import ServiceError
from app.services.transfer import free_name

log = logging.getLogger("datadesk.convertor")

DELIMITERS = {"comma": ",", "tab": "\t", "semicolon": ";", "pipe": "|"}
PREVIEW_ROWS = 5
CONVERTOR_FOLDER = ["Convertor"]


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_sheets(data: bytes, max_rows: int) -> list[Sheet]:
    """Every sheet of the workbook, as text rows."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ServiceError(f"Could not open the workbook: {exc}", status_code=422) from exc
    sheets = []
    try:
        for ws_ in wb.worksheets:
            rows: list[list[str]] = []
            for raw in ws_.iter_rows(values_only=True):
                cells = [cell_text(v) for v in raw]
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells:
                    continue  # a fully empty row
                rows.append(cells)
                if len(rows) > max_rows + 1:
                    raise ServiceError(
                        f"Sheet {ws_.title!r} has more than {max_rows:,} rows, "
                        "the most DataDesk keeps in a file.",
                        status_code=413,
                    )
            sheet = Sheet(ws_.title, rows)
            width = sheet.width
            sheet.rows = [r + [""] * (width - len(r)) for r in rows]
            sheets.append(sheet)
    finally:
        wb.close()
    return sheets


def to_csv(rows: list[list[str]], delimiter: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def csv_name(name: str) -> str:
    """A file name from what the person typed (or a sheet name): no slashes, ends in .csv."""
    clean = re.sub(r"[\\/:*?\"<>|]+", "-", name).strip().strip(".") or "Sheet"
    clean = clean[:180]
    return clean if clean.lower().endswith(".csv") else f"{clean}.csv"


async def _workbook(session: AsyncSession, user: User, file_id: uuid.UUID) -> tuple[WorkspaceFile, bytes]:
    file = await ws.owned_file(session, file_id, user.id)
    if file.format != "xlsx":
        if file.original_name.lower().endswith(".xls"):
            raise ServiceError("Old .xls files aren't supported: open it in Excel and save it as .xlsx.")
        raise ServiceError("Choose an Excel workbook (.xlsx).")
    return file, await ws.download(session, file, None)


async def inspect(session: AsyncSession, settings: Settings, user: User, file_id: uuid.UUID) -> dict:
    """The workbook's sheets: name, size and the first few rows."""
    file, data = await _workbook(session, user, file_id)
    sheets = await asyncio.to_thread(read_sheets, data, settings.max_rows)
    return {
        "file_id": file.id,
        "workbook": file.original_name,
        "folder_id": file.folder_id,
        "sheets": [
            {
                "name": s.name,
                "rows": max(len(s.rows) - 1, 0),  # the first row is usually the header
                "columns": s.width,
                "preview": s.rows[: PREVIEW_ROWS + 1],
            }
            for s in sheets
        ],
    }


async def convert(
    session: AsyncSession,
    settings: Settings,
    user: User,
    file_id: uuid.UUID,
    chosen: list[tuple[str, str]],
    delimiter: str,
    folder_id: uuid.UUID | None,
) -> tuple[list[WorkspaceFile], list[str]]:
    """Write each chosen (sheet, file name) as a CSV in the folder. Returns the new files and
    any "name → new name" for names that were already taken."""
    if not chosen:
        raise ServiceError("Choose at least one sheet.", field="sheets")
    sep = DELIMITERS.get(delimiter)
    if sep is None:
        raise ServiceError("Choose a delimiter: comma, tab, semicolon or pipe.", field="delimiter")
    await ws.folder_path(session, folder_id, user.id)  # it must be theirs

    file, data = await _workbook(session, user, file_id)
    sheets = {s.name: s for s in await asyncio.to_thread(read_sheets, data, settings.max_rows)}
    missing = [name for name, _ in chosen if name not in sheets]
    if missing:
        raise ServiceError(f"The workbook has no sheet named {', '.join(missing)}.", field="sheets")

    taken = set(
        await session.scalars(
            select(WorkspaceFile.original_name).where(
                WorkspaceFile.user_id == user.id,
                WorkspaceFile.status == "active",
                WorkspaceFile.folder_id.is_(None)
                if folder_id is None
                else WorkspaceFile.folder_id == folder_id,
            )
        )
    )
    created, renamed = [], []
    for sheet_name, wanted in chosen:
        sheet = sheets[sheet_name]
        if not sheet.rows:
            raise ServiceError(f"Sheet {sheet_name!r} is empty.", field="sheets")
        name = csv_name(wanted or sheet_name)
        final = free_name(name, taken, is_file=True)
        if final != name:
            renamed.append(f"{name} → {final}")
        taken.add(final)
        created.append(
            await ws.upload_file(session, settings, user, final, to_csv(sheet.rows, sep), folder_id)
        )
    log.info(
        "%s converted %d sheet(s) of %s to CSV (%s)",
        user.username,
        len(created),
        file.original_name,
        delimiter,
    )
    return created, renamed
